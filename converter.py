"""PDF to XLSX converter using Ollama vision models.

Renders each PDF page as an image, sends it to a local Ollama vision model
with a structured prompt describing the expected columns, and parses the
model's JSON response into spreadsheet rows.
"""

import base64
import io
import json
import os
import re
from pathlib import Path

import httpx
import pdfplumber
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from profiles import SupplierProfile

OLLAMA_BASE_URL = os.environ.get("OLLAMA_HOST", "http://localhost:11434")
DEFAULT_MODEL = os.environ.get("OLLAMA_MODEL", "qwen3-vl:235b-cloud")


def _render_page_to_base64(page, resolution: int = 150) -> str:
    """Render a pdfplumber page to a base64-encoded PNG string."""
    img = page.to_image(resolution=resolution).original
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("utf-8")


def _build_prompt(profile: SupplierProfile) -> str:
    """Build the extraction prompt for the vision model."""
    cols = " | ".join(profile.headers)
    num = len(profile.headers)

    prompt = f"""Extract the tabular data from this document image.

The table has {num} columns with these headers (in order):
{cols}

Rules:
- The table starts after the line containing "{profile.header_marker}".
"""
    if profile.stop_marker:
        prompt += f'- The table ends at the line containing "{profile.stop_marker}".\n'

    prompt += """- Return ONLY a JSON array of arrays. Each inner array is one row with exactly {num} string values.
- Preserve the exact order of columns.
- If a cell is empty, use an empty string "".
- Do NOT include the header row itself in the output.
- Do NOT include any text outside the JSON array.
- Numbers should be kept as strings exactly as they appear (e.g. "1.234,56" not "1234.56").
""".format(num=num)

    return prompt


def _extract_json_array(text: str) -> list[list[str]]:
    """Parse JSON array of arrays from model response, handling markdown fences."""
    text = text.strip()
    # Strip markdown code fences if present
    fence_match = re.search(r"```(?:json)?\s*\n?(.*?)```", text, re.DOTALL)
    if fence_match:
        text = fence_match.group(1).strip()

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        # Try to find the first [ ... ] block
        bracket_match = re.search(r"\[.*\]", text, re.DOTALL)
        if bracket_match:
            data = json.loads(bracket_match.group(0))
        else:
            return []

    if not isinstance(data, list):
        return []

    rows = []
    for item in data:
        if isinstance(item, list):
            rows.append([str(v) if v is not None else "" for v in item])
    return rows


def _call_ollama(image_b64: str, prompt: str, model: str) -> str:
    """Call Ollama vision API with an image and prompt."""
    with httpx.Client(timeout=600.0) as client:
        resp = client.post(
            f"{OLLAMA_BASE_URL}/api/chat",
            json={
                "model": model,
                "messages": [
                    {
                        "role": "user",
                        "content": prompt,
                        "images": [image_b64],
                    }
                ],
                "stream": False,
            },
        )
        resp.raise_for_status()
        return resp.json()["message"]["content"]


def check_ollama_available(model: str = DEFAULT_MODEL) -> tuple[bool, str]:
    """Check if Ollama is running and the model is accessible.

    Returns (ok, message) tuple.
    """
    try:
        with httpx.Client(timeout=15.0) as client:
            # First check Ollama is reachable
            resp = client.get(f"{OLLAMA_BASE_URL}/api/tags")
            resp.raise_for_status()

            # Try a lightweight call to verify the model works
            # (cloud models won't appear in local tags list)
            resp = client.post(
                f"{OLLAMA_BASE_URL}/api/chat",
                json={
                    "model": model,
                    "messages": [{"role": "user", "content": "hi"}],
                    "stream": False,
                },
                timeout=30.0,
            )
            if resp.status_code == 404:
                return False, (
                    f"Model '{model}' not found. "
                    f"Run: ollama pull {model}"
                )
            resp.raise_for_status()
            return True, f"OK — using {model}"
    except httpx.ConnectError:
        return False, (
            "Cannot connect to Ollama. "
            "Make sure Ollama is running (https://ollama.com)"
        )
    except Exception as e:
        return False, f"Ollama check failed: {e}"


def extract_rows_vision(
    pdf_path: str | Path,
    profile: SupplierProfile,
    model: str = DEFAULT_MODEL,
) -> list[list[str]]:
    """Extract table rows from PDF using Ollama vision model."""
    pdf_path = Path(pdf_path)
    prompt = _build_prompt(profile)
    num_cols = len(profile.headers)
    all_rows: list[list[str]] = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            image_b64 = _render_page_to_base64(page)
            response = _call_ollama(image_b64, prompt, model)
            rows = _extract_json_array(response)

            for row in rows:
                # Pad or trim to expected column count
                if len(row) < num_cols:
                    row.extend([""] * (num_cols - len(row)))
                elif len(row) > num_cols:
                    row = row[:num_cols]
                all_rows.append(row)

    return all_rows


def write_xlsx(
    headers: list[str],
    rows: list[list[str]],
    output_path: str | Path,
) -> Path:
    """Write headers + rows to an XLSX file with auto-filter and sizing."""
    output_path = Path(output_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    # Headers
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    # Data
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx).value = _maybe_number(value)

    # Auto-filter
    if rows:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    # Auto-size columns
    for col_idx, header in enumerate(headers, start=1):
        max_width = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_width = max(max_width, len(str(row[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 2, 50)

    if not rows:
        ws.cell(row=2, column=1, value="No matching rows found in this PDF.")

    wb.save(output_path)
    return output_path


def _maybe_number(value: str):
    """Try to convert a string to int or float."""
    if not value:
        return value
    cleaned = value.replace(" ", "").replace("\xa0", "")
    try:
        return int(cleaned)
    except ValueError:
        pass
    for attempt in [cleaned, cleaned.replace(",", ".")]:
        try:
            return float(attempt)
        except ValueError:
            pass
    return value


def convert_pdf_to_xlsx(
    pdf_path: str | Path,
    output_path: str | Path,
    profile: SupplierProfile,
) -> Path:
    """Main entry point: convert a PDF to XLSX using Ollama vision model."""
    pdf_path = Path(pdf_path)
    output_path = Path(output_path)

    rows = extract_rows_vision(pdf_path, profile)
    return write_xlsx(profile.headers, rows, output_path)
