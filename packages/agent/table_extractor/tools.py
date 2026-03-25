"""Tools for PDF table extraction and XLSX generation."""

import base64
import io
import json
import os
import re
from pathlib import Path
import pdfplumber
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


DEFAULT_GEMINI_MODEL = os.environ.get("GEMINI_MODEL", "gemini-2.0-flash")


# ---------------------------------------------------------------------------
# Vision model dispatch
# ---------------------------------------------------------------------------


async def _call_vision_model(
    image_b64: str,
    prompt: str,
    provider: str | None = None,
    model: str | None = None,
) -> str:
    """Call the vision model. Model can be overridden per-request."""
    return await _call_gemini(image_b64, prompt, model or DEFAULT_GEMINI_MODEL)


async def _call_gemini(image_b64: str, prompt: str, model: str) -> str:
    from google import genai
    from google.genai import types

    client = genai.Client()
    response = await client.aio.models.generate_content(
        model=model,
        contents=[
            types.Content(
                role="user",
                parts=[
                    types.Part.from_bytes(
                        data=base64.b64decode(image_b64),
                        mime_type="image/png",
                    ),
                    types.Part.from_text(text=prompt),
                ],
            )
        ],
    )
    return response.text


# ---------------------------------------------------------------------------
# PDF rendering
# ---------------------------------------------------------------------------

def _render_page_to_base64(page, resolution: int = 150) -> str:
    img = page.to_image(resolution=resolution).original
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("utf-8")


# ---------------------------------------------------------------------------
# JSON parsing
# ---------------------------------------------------------------------------

def _extract_json_array(text: str) -> list[list[str]]:
    text = text.strip()
    fence_match = re.search(r"```(?:json)?\s*\n?(.*?)```", text, re.DOTALL)
    if fence_match:
        text = fence_match.group(1).strip()

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        bracket_match = re.search(r"\[.*\]", text, re.DOTALL)
        if bracket_match:
            data = json.loads(bracket_match.group(0))
        else:
            return []

    if not isinstance(data, list):
        return []

    rows = []
    for item in data:
        if isinstance(item, list):
            rows.append([str(v) if v is not None else "" for v in item])
    return rows


# ---------------------------------------------------------------------------
# XLSX generation
# ---------------------------------------------------------------------------

def _maybe_number(value: str):
    if not value:
        return value
    cleaned = value.replace(" ", "").replace("\xa0", "")
    try:
        return int(cleaned)
    except ValueError:
        pass
    for attempt in [cleaned, cleaned.replace(",", ".")]:
        try:
            return float(attempt)
        except ValueError:
            pass
    return value


def _write_xlsx(
    headers: list[str],
    rows: list[list[str]],
    output_path: str | Path,
) -> Path:
    output_path = Path(output_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx).value = _maybe_number(value)

    if rows:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    for col_idx, header in enumerate(headers, start=1):
        max_width = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_width = max(max_width, len(str(row[col_idx - 1])))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width + 2, 50)

    if not rows:
        ws.cell(row=2, column=1, value="No matching rows found in this PDF.")

    wb.save(output_path)
    return output_path


